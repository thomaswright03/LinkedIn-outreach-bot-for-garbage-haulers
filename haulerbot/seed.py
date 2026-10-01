"""Import the Soft-Pak prospect seed workbook (softpak_mts_prospect_seed.xlsx) or a CSV.

Two layouts are accepted, and column names are matched loosely (case, punctuation and
common variants don't matter):

* one row per company, with decision-makers in role columns such as "Owner/CEO",
  "President", "COO", "Sales/Growth Executive";
* one row per contact, with "Contact Name" / "Title" / "LinkedIn" / "Email" columns
  alongside the company columns.

Emails are stored as verified only when a verification column explicitly says so.
An email that is only a guess from the company's address pattern stays unverified.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

from .db import DB, PRIORITIES
from .discovery import canonical_profile_url, first_name_of


def _norm(header) -> str:
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower())


# normalized header -> company field
COMPANY_COLUMNS = {
    "company": "name", "companyname": "name", "hauler": "name", "haulercompany": "name",
    "haulername": "name", "account": "name", "accountname": "name",
    "website": "website", "companywebsite": "website", "domain": "website",
    "headquarters": "headquarters", "hq": "headquarters", "hqlocation": "headquarters",
    "location": "headquarters", "city": "headquarters",
    "statesmarkets": "markets", "markets": "markets", "states": "markets",
    "marketsstates": "markets", "servicearea": "markets", "state": "markets",
    "residentialservice": "residential_service", "residential": "residential_service",
    "estimatedresidentialaccounts": "est_residential_accounts",
    "estresidentialaccounts": "est_residential_accounts",
    "residentialaccounts": "est_residential_accounts",
    "softpakstatus": "softpak_status", "softpak": "softpak_status",
    "softpakconfidence": "softpak_status",
    "softpakevidence": "softpak_evidence", "evidence": "softpak_evidence",
    "evidenceurl": "evidence_url", "source": "evidence_url", "sourceurl": "evidence_url",
    "evidencelink": "evidence_url",
    "businessphone": "business_phone", "phone": "business_phone", "mainphone": "business_phone",
    "companylinkedin": "linkedin_url", "linkedincompanypage": "linkedin_url",
    "mtsfit": "mts_fit", "fit": "mts_fit", "fitnotes": "mts_fit",
    "priority": "priority", "tier": "priority", "mtspriority": "priority",
    "nextaction": "next_action", "notes": "notes", "researchnotes": "notes",
    "demobooked": "demo_booked", "demodate": "demo_date",
}

# normalized header -> role label, for the one-row-per-company layout
ROLE_COLUMNS = {
    "ownerceo": "Owner/CEO", "owner": "Owner/CEO", "ceo": "Owner/CEO",
    "president": "President", "coo": "COO",
    "gmregionalpresident": "GM/Regional President", "gm": "GM/Regional President",
    "generalmanager": "GM/Regional President", "regionalpresident": "GM/Regional President",
    "salesgrowthexecutive": "Sales/Growth Executive", "salesexecutive": "Sales/Growth Executive",
    "vpsales": "Sales/Growth Executive", "salesgrowth": "Sales/Growth Executive",
    "marketingexecutive": "Marketing Executive", "marketing": "Marketing Executive",
    "itcio": "IT/CIO", "cio": "IT/CIO", "it": "IT/CIO",
}

# normalized header -> contact field, for the one-row-per-contact layout
CONTACT_COLUMNS = {
    "contactname": "full_name", "contact": "full_name", "fullname": "full_name",
    "decisionmaker": "full_name",
    "title": "role", "role": "role", "contacttitle": "role", "jobtitle": "role",
    "linkedin": "linkedin", "linkedinurl": "linkedin", "linkedinprofile": "linkedin",
    "contactlinkedin": "linkedin",
    "email": "email", "businessemail": "email", "contactemail": "email",
    "emailverified": "email_verified", "emailstatus": "email_verified",
    "verified": "email_verified", "emailverification": "email_verified",
    "contactphone": "phone", "directphone": "phone", "mobile": "phone",
}

_PLACEHOLDER = re.compile(r"^(tbd|tba|unknown|n/?a|none|-+|\?+|todo|research|not found)$", re.I)
_URL = re.compile(r"https?://\S+")


@dataclass
class ImportReport:
    companies: int = 0
    contacts: int = 0
    contacts_with_linkedin: int = 0
    skipped_rows: int = 0
    sheets: list[str] = field(default_factory=list)
    unverified_emails: int = 0


def _clean(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return "" if _PLACEHOLDER.match(s) else s


def normalize_priority(v: str) -> str:
    m = re.match(r"\s*(?:tier\s*|priority\s*)?([abc])\b", v or "", re.I)
    return m.group(1).upper() if m and m.group(1).upper() in PRIORITIES else ""


def normalize_softpak(v: str) -> str:
    t = (v or "").lower()
    if not t:
        return ""
    if "histor" in t or "former" in t or "past" in t:
        return "historical"
    if "confirm" in t or t in ("yes", "y", "verified"):
        return "confirmed"
    if any(w in t for w in ("probable", "likely", "high", "possible", "suspected")):
        return "probable"
    return ""


def is_verified_flag(v: str) -> bool:
    """Only an explicit 'verified'-style value counts. 'Guessed', 'pattern', blank: no."""
    t = (v or "").strip().lower()
    if not t or any(w in t for w in ("guess", "pattern", "unverified", "not verified", "inferred")):
        return False
    return t in ("yes", "y", "true", "1", "verified", "valid", "deliverable", "confirmed")


def _split_people(cell: str) -> list[tuple[str, str]]:
    """'Jane Doe (CEO) https://linkedin.com/in/jd; Bob Ray' -> [(name, url), ...]."""
    out = []
    for part in re.split(r"[;\n]|\s/\s", cell):
        part = part.strip()
        if not part:
            continue
        url = ""
        m = _URL.search(part)
        if m:
            url = m.group(0)
            part = (part[:m.start()] + part[m.end():]).strip()
        name = re.split(r"\s+[-–—|]\s+|\(|,", part)[0].strip()
        if name and not _PLACEHOLDER.match(name) and re.search(r"[A-Za-z]", name):
            out.append((name, url))
    return out


def _rows_from_csv(path: Path):
    with path.open(newline="", encoding="utf-8-sig") as f:
        yield path.name, list(csv.reader(f))


def _rows_from_xlsx(path: Path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        yield ws.title, [list(r) for r in ws.iter_rows(values_only=True)]


def _header_index(rows: list[list]) -> int | None:
    for i, row in enumerate(rows[:15]):
        cells = {_norm(c) for c in row}
        if cells & {k for k, v in COMPANY_COLUMNS.items() if v == "name"}:
            return i
    return None


def import_seed(db: DB, path: str | Path) -> ImportReport:
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        sheets = _rows_from_xlsx(path)
    elif path.suffix.lower() == ".csv":
        sheets = _rows_from_csv(path)
    else:
        raise ValueError("seed file must be .xlsx or .csv")

    report = ImportReport()
    seen_companies: set[int] = set()
    for title, rows in sheets:
        hi = _header_index(rows)
        if hi is None:
            continue
        report.sheets.append(title)
        headers = [_norm(h) for h in rows[hi]]
        company_cols = {i: COMPANY_COLUMNS[h] for i, h in enumerate(headers) if h in COMPANY_COLUMNS}
        role_cols = {i: ROLE_COLUMNS[h] for i, h in enumerate(headers) if h in ROLE_COLUMNS}
        contact_cols = {i: CONTACT_COLUMNS[h] for i, h in enumerate(headers) if h in CONTACT_COLUMNS}
        per_contact = "full_name" in contact_cols.values()

        for row in rows[hi + 1:]:
            vals = {i: _clean(row[i]) if i < len(row) else "" for i in range(len(headers))}
            fields = {f: vals[i] for i, f in company_cols.items() if vals[i]}
            name = fields.pop("name", "")
            if not name:
                report.skipped_rows += 1
                continue
            if "priority" in fields:
                fields["priority"] = normalize_priority(fields["priority"])
            if "softpak_status" in fields:
                fields["softpak_status"] = normalize_softpak(fields["softpak_status"])
            demo_booked = fields.pop("demo_booked", "")
            demo_date = fields.pop("demo_date", "")
            if fields.get("linkedin_url") and "/company/" not in fields["linkedin_url"]:
                fields.pop("linkedin_url")
            cid = db.upsert_company(name, **fields)
            if demo_booked.lower() in ("yes", "y", "true", "1") or demo_date:
                db.conn.execute("UPDATE companies SET demo_booked = 1, demo_date = ? WHERE id = ?",
                                (demo_date, cid))
                db.conn.commit()
            if cid not in seen_companies:
                seen_companies.add(cid)
                report.companies += 1

            contact = {f: vals[i] for i, f in contact_cols.items() if vals[i]}
            if per_contact:
                people = [(contact.get("full_name", ""), contact.get("linkedin", ""),
                           contact.get("role", ""))]
            else:
                people = [(n, u, role) for i, role in role_cols.items()
                          for n, u in _split_people(vals[i])]
                # A single LinkedIn/email column only belongs to a contact if there's one.
                if len(people) == 1 and contact.get("linkedin") and not people[0][1]:
                    n, _, role = people[0]
                    people = [(n, contact["linkedin"], role)]
            single = len(people) == 1
            for full_name, url, role in people:
                if not full_name:
                    continue
                email = contact.get("email", "") if single else ""
                verified = single and is_verified_flag(contact.get("email_verified", ""))
                profile = canonical_profile_url(url) if url else None
                pid = db.add_contact(
                    cid, full_name, first_name_of(full_name), role=role, profile_url=profile,
                    email=email, email_verified=verified,
                    phone=contact.get("phone", "") if single else "")
                if pid:
                    report.contacts += 1
                    report.contacts_with_linkedin += bool(profile)
                    report.unverified_emails += bool(email and not verified)
    return report
